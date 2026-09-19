import os
import json
import openpyxl
import psycopg2
from psycopg2.extras import execute_values
from dotenv import load_dotenv
import pathlib

load_dotenv()

conn = psycopg2.connect(os.environ.get('DATABASE_URL'))
cursor = conn.cursor()

SCRIPT_DIR = pathlib.Path(__file__).resolve().parent
file_path = SCRIPT_DIR / "Data" / "recipes_with_vectors.xlsx"
wb = openpyxl.load_workbook(file_path, read_only=True)
ws = wb['Recipes']

rows_to_insert = []
for row in ws.iter_rows(min_row=2, values_only=True):
    recipe_id, name, description, servings, prep_time, cook_time, vector_str = row

    if vector_str is None:
        continue

    vector_list = [float(x) for x in vector_str.split(',')]
    vector_json = json.dumps(vector_list)

    rows_to_insert.append((recipe_id, name, description, servings, prep_time, cook_time, vector_json))

print(f"Loaded {len(rows_to_insert)} recipes from spreadsheet. Inserting...")

execute_values(
    cursor,
    """
    INSERT INTO recipes (recipe_id, name, description, servings, prep_time, cook_time, recipe_vector)
    VALUES %s
    ON CONFLICT (recipe_id) DO NOTHING
    """,
    rows_to_insert
)

print("Loading directions and ingredients...")

wb2 = openpyxl.load_workbook(SCRIPT_DIR / "Data" / "recipe_database.xlsx", read_only=True)

directions_rows = []
for row in wb2['Directions'].iter_rows(min_row=2, values_only=True):
    directions_rows.append(row)

execute_values(
    cursor,
    "INSERT INTO directions (step_id, recipe_id, instruction_text) VALUES %s ON CONFLICT (step_id) DO NOTHING",
    directions_rows
)

ingredients_rows = []
for row in wb2['Ingredients'].iter_rows(min_row=2, values_only=True):
    ingredients_rows.append(row)

execute_values(
    cursor,
    "INSERT INTO ingredients (ingredient_id, recipe_id, ingredient_text) VALUES %s ON CONFLICT (ingredient_id) DO NOTHING",
    ingredients_rows
)

conn.commit()
print(f"Loaded {len(directions_rows)} direction steps and {len(ingredients_rows)} ingredients.")

conn.commit()
cursor.close()
conn.close()

print("Done.")